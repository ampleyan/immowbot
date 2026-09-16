import asyncio
import hashlib
import hmac
import io
import json
import os
import queue
import re
import sqlite3
import sys
import threading
import time
from datetime import datetime, timezone

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..")))

from fastapi import FastAPI, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from sse_starlette.sse import EventSourceResponse

from src.buyer.property_scoring import calculate_home_score
from src.buyer.purchase_calculator import calculate_purchase_estimate
from src.buyer.property_store import PropertyStore
from src.buyer.search_config import DEFAULT_HOME_SEARCH, normalize_search_config
from src.buyer.smart_lists import BUILTIN_SMART_LISTS, explain_rule_match, matches_rule
from src.buyer.change_tracking import diff_versions
from src.buyer.property_explanation import explain_property
from src.buyer.commute import commute_estimate
from src.buyer.duplicate_detection import duplicate_groups

DB_PATH = os.getenv("DB_PATH", os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "data", "buyer.db")))
SEARCH_NAME = "antwerp-home"

PHONE_RE = re.compile(r'(?:\+32|0032|0)\s*\d[\d\s.\-/]{6,12}\d')
EMAIL_RE = re.compile(r'[\w.+-]+@[\w-]+\.[a-z]{2,}', re.IGNORECASE)

def extract_contact_from_listing(listing):
    text = ' '.join(filter(None, [
        listing.get('description') or '',
        listing.get('description_english') or '',
        str((listing.get('all_property_details') or {}).get('Description (Original)') or ''),
        str((listing.get('all_property_details') or {}).get('Description (English)') or ''),
    ]))
    phone = PHONE_RE.search(text)
    email = EMAIL_RE.search(text)
    return {
        'phone': re.sub(r'\s+', ' ', phone.group()).strip() if phone else None,
        'email': email.group().strip() if email else None,
    }

app = FastAPI()
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

AUTH_SECRET = os.getenv("IMMOWBOT_AUTH_SECRET", "change-me-in-production")
AUTH_COOKIE = "immowbot_session"
INVITE_TOKEN = os.getenv("IMMOWBOT_INVITE_TOKEN", "")
TRUSTED_IPS = set(filter(None, os.getenv("TRUSTED_IPS", "127.0.0.1,::1").split(",")))


def _get_client_ip(request):
    forwarded = request.headers.get("X-Forwarded-For")
    if forwarded:
        return forwarded.split(",")[0].strip()
    return request.client.host if request.client else None


def _is_trusted(request):
    return bool(TRUSTED_IPS) and _get_client_ip(request) in TRUSTED_IPS


def _session_token(username):
    payload = f"{username}:{int(time.time()) // 86400}"
    signature = hmac.new(AUTH_SECRET.encode(), payload.encode(), hashlib.sha256).hexdigest()
    return f"{payload}:{signature}"


def _session_username(token):
    if not token:
        return None
    try:
        username, day, signature = token.split(":", 2)
        payload = f"{username}:{day}"
        expected = hmac.new(AUTH_SECRET.encode(), payload.encode(), hashlib.sha256).hexdigest()
        if int(day) >= int(time.time()) // 86400 - 1 and hmac.compare_digest(signature, expected):
            return username
        return None
    except (ValueError, TypeError):
        return None


def get_store():
    print(f"[immowbot] Using DB: {DB_PATH}", flush=True)
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    return PropertyStore(DB_PATH)


def _get_current_user(request):
    username = _session_username(request.cookies.get(AUTH_COOKIE))
    if not username and _is_trusted(request):
        username = "ampleyan"
    if not username:
        return None
    store = get_store()
    try:
        return store.get_user_by_username(username)
    finally:
        store.close()


def _require_current_user(request):
    user = _get_current_user(request)
    if not user:
        raise HTTPException(401, "Authentication required")
    return user


def _require_admin(request):
    user = _require_current_user(request)
    if not user["is_admin"]:
        raise HTTPException(403, "Admin access required")
    return user


@app.middleware("http")
async def require_login(request: Request, call_next):
    public = {"/api/auth/login", "/api/auth/login-trusted", "/api/auth/trusted", "/api/auth/me", "/api/health"}
    if request.url.path.startswith("/api/") and request.url.path not in public and not request.url.path.startswith("/api/register/"):
        if not _session_username(request.cookies.get(AUTH_COOKIE)) and not _is_trusted(request):
            return JSONResponse({"detail": "Authentication required"}, status_code=401)
    return await call_next(request)


APP_VERSION = os.getenv("APP_VERSION", "")
if not APP_VERSION or APP_VERSION == "dev":
    try:
        with open(os.path.join(os.path.dirname(__file__), "../../VERSION")) as f:
            APP_VERSION = f.read().strip()
    except Exception:
        APP_VERSION = "dev"


@app.get("/api/health")
async def health():
    return {"status": "ok", "version": APP_VERSION}


@app.post("/api/auth/login")
async def login(body: dict):
    username = str(body.get("username") or "")
    password = str(body.get("password") or "")
    store = get_store()
    try:
        user = store.authenticate_user(username, password)
    finally:
        store.close()
    if not user:
        raise HTTPException(401, "Invalid username or password")
    response = JSONResponse({"username": user["username"], "is_admin": bool(user["is_admin"])})
    response.set_cookie(AUTH_COOKIE, _session_token(user["username"]), httponly=True, samesite="lax", secure=False, max_age=172800)
    return response


@app.get("/api/auth/me")
async def auth_me(request: Request):
    user = _get_current_user(request)
    if not user:
        raise HTTPException(401, "Authentication required")
    return {"username": user["username"], "is_admin": bool(user["is_admin"])}


@app.post("/api/auth/logout")
async def logout():
    response = JSONResponse({"ok": True})
    response.delete_cookie(AUTH_COOKIE)
    return response


@app.get("/api/auth/trusted")
async def check_trusted(request: Request):
    return {"trusted": _is_trusted(request)}


@app.post("/api/auth/login-trusted")
async def login_trusted(request: Request):
    if not _is_trusted(request):
        raise HTTPException(403, "Not a trusted IP address")
    store = get_store()
    try:
        user = store.get_user_by_username("ampleyan")
    finally:
        store.close()
    if not user:
        raise HTTPException(404, "User not found")
    response = JSONResponse({"username": user["username"], "is_admin": bool(user["is_admin"])})
    response.set_cookie(AUTH_COOKIE, _session_token(user["username"]), httponly=True, samesite="lax", secure=False, max_age=172800)
    return response


@app.get("/api/register/{token}")
async def check_invite(token: str):
    if not INVITE_TOKEN or not hmac.compare_digest(token, INVITE_TOKEN):
        raise HTTPException(403, "Invalid invite link")
    return {"valid": True}


@app.post("/api/register/{token}")
async def register(token: str, body: dict):
    if not INVITE_TOKEN or not hmac.compare_digest(token, INVITE_TOKEN):
        raise HTTPException(403, "Invalid invite link")
    username = str(body.get("username") or "").strip()
    password = str(body.get("password") or "")
    if not username or len(username) < 3 or " " in username:
        raise HTTPException(400, "Username must be at least 3 characters with no spaces")
    if len(password) < 8:
        raise HTTPException(400, "Password must be at least 8 characters")
    store = get_store()
    try:
        try:
            user_id = store.create_user(username, password, is_admin=False)
        except Exception:
            raise HTTPException(409, "Username already taken")
        user = store.get_user_by_id(user_id)
    finally:
        store.close()
    response = JSONResponse({"username": user["username"], "is_admin": False})
    response.set_cookie(AUTH_COOKIE, _session_token(user["username"]), httponly=True, samesite="lax", secure=False, max_age=172800)
    return response


_state = {"search_ids": {}, "collection": None, "translation": None}

BACKUP_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "data", "backups"))


def save_listings_backup(store, run_id):
    os.makedirs(BACKUP_DIR, exist_ok=True)
    listings = store.latest_listings("sale")
    ts = datetime.now(timezone.utc).strftime("%Y-%m-%d_%H-%M-%S")
    path = os.path.join(BACKUP_DIR, f"{ts}_run-{run_id}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(listings, f, ensure_ascii=False, indent=2, default=str)


def _translation_state():
    t = _state.get("translation")
    if not t:
        return {"translating": False}
    return {
        "translating": True,
        "translation_done": t["done"],
        "translation_total": t["total"],
        "translation_current": t.get("current", ""),
        "translation_current_address": t.get("current_address", ""),
        "translation_current_score": t.get("current_score", None),
    }


def _get_or_init_search_id(store, user_id):
    if user_id in _state["search_ids"]:
        return _state["search_ids"][user_id]
    existing = store.get_search_by_name(user_id, SEARCH_NAME)
    if existing:
        config = {**DEFAULT_HOME_SEARCH, **existing["config"]}
        sid = store.save_search(user_id, SEARCH_NAME, existing["purpose"], config)
    else:
        sid = store.save_search(user_id, SEARCH_NAME, "home", DEFAULT_HOME_SEARCH)
    store.ensure_builtin_smart_lists(user_id, BUILTIN_SMART_LISTS)
    _state["search_ids"][user_id] = sid
    return sid


@app.on_event("startup")
def startup():
    store = get_store()
    try:
        for user in store.list_users():
            _get_or_init_search_id(store, user["id"])
    finally:
        store.close()


@app.get("/api/users")
def list_users(request: Request):
    _require_admin(request)
    store = get_store()
    try:
        return store.list_users()
    finally:
        store.close()


@app.post("/api/users")
def create_user(request: Request, body: dict):
    _require_admin(request)
    username = str(body.get("username") or "").strip()
    password = str(body.get("password") or "")
    is_admin = bool(body.get("is_admin", False))
    if not username or not password:
        raise HTTPException(400, "username and password required")
    store = get_store()
    try:
        try:
            user_id = store.create_user(username, password, is_admin)
        except sqlite3.IntegrityError:
            raise HTTPException(409, "username already exists")
        _get_or_init_search_id(store, user_id)
        return {"id": user_id, "username": username, "is_admin": is_admin}
    finally:
        store.close()


@app.put("/api/users/me/password")
def change_password(request: Request, body: dict):
    user = _require_current_user(request)
    current_password = str(body.get("current_password") or "")
    new_password = str(body.get("new_password") or "")
    if len(new_password) < 8:
        raise HTTPException(400, "New password must be at least 8 characters")
    store = get_store()
    try:
        if not store.authenticate_user(user["username"], current_password):
            raise HTTPException(403, "Current password is incorrect")
        store.update_user_password(user["id"], new_password)
        return {"ok": True}
    finally:
        store.close()


@app.delete("/api/users/{target_user_id}")
def delete_user(request: Request, target_user_id: int):
    current = _require_admin(request)
    if current["id"] == target_user_id:
        raise HTTPException(400, "cannot delete your own account")
    store = get_store()
    try:
        store.delete_user(target_user_id)
        _state["search_ids"].pop(target_user_id, None)
        return {"ok": True}
    finally:
        store.close()


@app.get("/api/listings")
def get_listings(request: Request):
    user = _require_current_user(request)
    user_id = user["id"]
    store = get_store()
    try:
        search_id = _get_or_init_search_id(store, user_id)
        config = store.get_search(search_id)["config"]
        listings = store.latest_listings("sale")
        postcode_avgs = {}
        for listing in listings:
            pc = listing.get("postcode")
            price = listing.get("price")
            surface = listing.get("surface_area")
            if pc and price and surface and surface > 0:
                postcode_avgs.setdefault(pc, []).append(price / surface)
        postcode_avg_price_per_sqm = {pc: sum(vals) / len(vals) for pc, vals in postcode_avgs.items()}
        for listing in listings:
            listing["_postcode_avg_price_per_sqm"] = postcode_avg_price_per_sqm.get(listing.get("postcode"))
        duplicate_keys = {
            (offer.get("source"), str(offer.get("source_listing_id", "")))
            for group in duplicate_groups(listings)
            for offer in group["offers"]
        }
        all_notes = store.get_all_notes(user_id)
        result = []
        for listing in listings:
            scored = calculate_home_score(listing, config)
            src = listing.get("source", "")
            lid = str(listing.get("source_listing_id", ""))
            list_ids = list(store.get_property_list_ids(user_id, src, lid))
            note = all_notes.get((src, lid), "")
            if not listing.get("agent_phone") or not listing.get("agent_email"):
                extracted = extract_contact_from_listing(listing)
                if not listing.get("agent_phone") and extracted["phone"]:
                    listing["agent_phone"] = extracted["phone"]
                if not listing.get("agent_email") and extracted["email"]:
                    listing["agent_email"] = extracted["email"]
            result.append({
                **listing,
                "_is_duplicate": (src, lid) in duplicate_keys,
                "_score": scored["score"],
                "_components": scored["components"],
                "_score_weights": config.get("score_weights"),
                "_exclusions": scored["exclusions"],
                "_list_ids": list_ids,
                "_note": note,
                "_purchase_estimate": calculate_purchase_estimate(listing, config),
                "_workflow": store.get_workflow(user_id, src, lid),
                "_explanation": explain_property(listing, scored["score"], scored["components"], scored["exclusions"], calculate_purchase_estimate(listing, config)),
                "_commute": commute_estimate(listing, config.get("commute_destinations", [])),
            })
        result.sort(key=lambda x: (x["_score"] is None, -(x["_score"] or 0)))
        return result
    finally:
        store.close()


@app.post("/api/listings/batch")
def get_listings_batch(request: Request, body: dict):
    user = _require_current_user(request)
    user_id = user["id"]
    store = get_store()
    try:
        ids = set()
        for item in (body.get("ids") or []):
            if item.get("source") and item.get("source_listing_id"):
                ids.add((item["source"], str(item["source_listing_id"])))
        if not ids:
            return []
        search_id = _get_or_init_search_id(store, user_id)
        config = store.get_search(search_id)["config"]
        all_listings = store.latest_listings("sale")
        postcode_avgs = {}
        for listing in all_listings:
            pc = listing.get("postcode")
            price = listing.get("price")
            surface = listing.get("surface_area")
            if pc and price and surface and surface > 0:
                postcode_avgs.setdefault(pc, []).append(price / surface)
        postcode_avg_price_per_sqm = {pc: sum(vals) / len(vals) for pc, vals in postcode_avgs.items()}
        for listing in all_listings:
            listing["_postcode_avg_price_per_sqm"] = postcode_avg_price_per_sqm.get(listing.get("postcode"))
        listings = [l for l in all_listings if (l.get("source"), str(l.get("source_listing_id", ""))) in ids]
        all_notes = store.get_all_notes(user_id)
        result = []
        for listing in listings:
            scored = calculate_home_score(listing, config)
            score = scored["score"]
            if score is None and scored.get("components"):
                score = min(100, round(sum(scored["components"].values()), 2))
            src = listing.get("source", "")
            lid = str(listing.get("source_listing_id", ""))
            result.append({
                **listing,
                "_score": score,
                "_components": scored["components"],
                "_score_weights": config.get("score_weights"),
                "_exclusions": scored["exclusions"],
                "_list_ids": list(store.get_property_list_ids(user_id, src, lid)),
                "_note": all_notes.get((src, lid), ""),
                "_workflow": store.get_workflow(user_id, src, lid),
                "_explanation": explain_property(listing, scored["score"], scored["components"], scored["exclusions"], calculate_purchase_estimate(listing, config)),
                "_commute": commute_estimate(listing, config.get("commute_destinations", [])),
            })
        return result
    finally:
        store.close()


@app.get("/api/config")
def get_config(request: Request):
    user = _require_current_user(request)
    store = get_store()
    try:
        search_id = _get_or_init_search_id(store, user["id"])
        return store.get_search(search_id)["config"]
    finally:
        store.close()


@app.put("/api/config")
def update_config(request: Request, body: dict):
    user = _require_current_user(request)
    user_id = user["id"]
    store = get_store()
    try:
        search_id = _get_or_init_search_id(store, user_id)
        current = store.get_search(search_id)["config"]
        merged = {**current, **body}
        try:
            normalize_search_config(merged)
        except ValueError as exc:
            raise HTTPException(400, str(exc))
        new_id = store.save_search(user_id, SEARCH_NAME, "home", merged)
        _state["search_ids"][user_id] = new_id
        return store.get_search(new_id)["config"]
    finally:
        store.close()


@app.get("/api/runs")
def get_runs(request: Request):
    user = _require_current_user(request)
    store = get_store()
    try:
        search_id = _get_or_init_search_id(store, user["id"])
        rows = store.connection.execute(
            """SELECT r.id, r.started_at, r.completed_at, r.status,
                      GROUP_CONCAT(sr.source || ':' || sr.status || ':' || sr.listing_count, '|') as sources_raw
               FROM runs r
               LEFT JOIN source_runs sr ON sr.run_id = r.id
               WHERE r.search_id = ?
               GROUP BY r.id
               ORDER BY r.id DESC
               LIMIT 20""",
            (search_id,),
        ).fetchall()
        runs = []
        for row in rows:
            sources = []
            if row["sources_raw"]:
                for part in row["sources_raw"].split("|"):
                    bits = part.split(":")
                    if len(bits) == 3:
                        sources.append({"source": bits[0], "status": bits[1], "count": int(bits[2])})
            runs.append({
                "id": row["id"],
                "started_at": row["started_at"],
                "completed_at": row["completed_at"],
                "status": row["status"],
                "sources": sources,
            })
        return runs
    finally:
        store.close()


@app.get("/api/runs/stream")
async def stream_progress(request: Request):
    _require_current_user(request)

    async def generator():
        while True:
            col = _state.get("collection")
            if col and col["thread"].is_alive():
                while True:
                    try:
                        col["progress"].update(col["progress_queue"].get_nowait())
                    except queue.Empty:
                        break
                p = col["progress"]
                yield {"data": json.dumps({
                    "alive": True,
                    "checked": p.get("checked", 0),
                    "saved": p.get("saved", 0),
                    "selected": col.get("selected", 0),
                    "portal": p.get("portal", ""),
                    "cancelling": col["cancel_event"].is_set(),
                    **_translation_state(),
                })}
            else:
                if col:
                    while True:
                        try:
                            col["progress"].update(col["progress_queue"].get_nowait())
                        except queue.Empty:
                            break
                    p = col["progress"]
                    _state["collection"] = None
                    yield {"data": json.dumps({
                        "alive": False,
                        "status": p.get("status", "ok"),
                        "checked": p.get("checked", 0),
                        "saved": p.get("saved", 0),
                        "selected": col.get("selected", 0),
                        "error": p.get("error"),
                        **_translation_state(),
                    })}
                else:
                    yield {"data": json.dumps({"alive": False, **_translation_state()})}
            await asyncio.sleep(0.8)

    return EventSourceResponse(generator())


@app.delete("/api/runs/active")
def cancel_run(request: Request):
    _require_current_user(request)
    col = _state.get("collection")
    if col and col["thread"].is_alive():
        col["cancel_event"].set()
        return {"ok": True}
    return {"ok": False}


@app.post("/api/runs")
def start_run(request: Request):
    user = _require_current_user(request)
    user_id = user["id"]
    if _state.get("collection") and _state["collection"]["thread"].is_alive():
        raise HTTPException(409, "Collection already running")
    store = get_store()
    try:
        search_id = _get_or_init_search_id(store, user_id)
    finally:
        store.close()

    def _worker(store_path, sid, cancel_event, progress_queue):
        from src.buyer.collector import run_collection
        from src.scraper_manager import ScraperManager
        s = PropertyStore(store_path)
        manager = ScraperManager()

        def on_translate(done, total, current):
            if done >= total:
                _state["translation"] = None
            else:
                _state["translation"] = {"total": total, "done": done, "current": current}

        try:
            run_id = run_collection(s, sid, manager, on_progress=progress_queue.put, should_cancel=cancel_event.is_set, on_translate_progress=on_translate)
            progress_queue.put({"status": s.get_run(run_id)["status"]})
            save_listings_backup(s, run_id)
        except Exception as exc:
            progress_queue.put({"status": "error", "error": str(exc)})
        finally:
            _state["translation"] = None
            s.close()

    cancel_event = threading.Event()
    progress_queue = queue.Queue()
    thread = threading.Thread(target=_worker, args=(DB_PATH, search_id, cancel_event, progress_queue), daemon=True)
    _state["collection"] = {
        "thread": thread,
        "cancel_event": cancel_event,
        "progress_queue": progress_queue,
        "selected": 0,
        "progress": {"checked": 0, "saved": 0, "status": "running"},
    }
    thread.start()
    return {"ok": True}


@app.post("/api/runs/selected")
def start_selected_run(request: Request, body: dict):
    user = _require_current_user(request)
    user_id = user["id"]
    if _state.get("collection") and _state["collection"]["thread"].is_alive():
        raise HTTPException(409, "Collection already running")
    requested = body.get("listings")
    if not isinstance(requested, list) or not requested or len(requested) > 100:
        raise HTTPException(400, "one to one hundred listings are required")
    store = get_store()
    try:
        search_id = _get_or_init_search_id(store, user_id)
        available = {
            (listing.get("source"), str(listing.get("source_listing_id"))): listing
            for listing in store.latest_listings("sale")
        }
        selections = []
        seen = set()
        for item in requested:
            if not isinstance(item, dict):
                continue
            key = (item.get("source"), str(item.get("source_listing_id", "")))
            listing = available.get(key)
            if listing and key not in seen and listing.get("url"):
                selections.append({"source": key[0], "source_listing_id": key[1], "url": listing["url"]})
                seen.add(key)
        if not selections:
            raise HTTPException(400, "no valid stored listings were selected")
    finally:
        store.close()

    def _worker(store_path, sid, selected, cancel_event, progress_queue):
        from src.buyer.collector import run_selected_collection
        from src.scraper_manager import ScraperManager
        s = PropertyStore(store_path)

        def on_translate(done, total, current):
            if done >= total:
                _state["translation"] = None
            else:
                _state["translation"] = {"total": total, "done": done, "current": current}

        try:
            run_id = run_selected_collection(s, sid, ScraperManager(), selected, on_progress=progress_queue.put, should_cancel=cancel_event.is_set, on_translate_progress=on_translate)
            progress_queue.put({"status": s.get_run(run_id)["status"]})
            save_listings_backup(s, run_id)
        except Exception as exc:
            progress_queue.put({"status": "error", "error": str(exc)})
        finally:
            _state["translation"] = None
            s.close()

    cancel_event = threading.Event()
    progress_queue = queue.Queue()
    thread = threading.Thread(target=_worker, args=(DB_PATH, search_id, selections, cancel_event, progress_queue), daemon=True)
    _state["collection"] = {
        "thread": thread,
        "cancel_event": cancel_event,
        "progress_queue": progress_queue,
        "selected": len(selections),
        "progress": {"checked": 0, "saved": 0, "status": "running"},
    }
    thread.start()
    return {"ok": True, "selected": len(selections)}


@app.get("/api/runs/{run_id}/listings")
def get_run_listings(request: Request, run_id: int):
    _require_current_user(request)
    store = get_store()
    try:
        return store.listings_for_run(run_id)
    finally:
        store.close()


@app.get("/api/lists")
def get_lists(request: Request):
    user = _require_current_user(request)
    store = get_store()
    try:
        return store.get_lists(user["id"])
    finally:
        store.close()


@app.post("/api/lists")
def create_list(request: Request, body: dict):
    user = _require_current_user(request)
    name = (body.get("name") or "").strip()
    if not name:
        raise HTTPException(400, "name required")
    store = get_store()
    try:
        list_id = store.create_list(user["id"], name)
        return {"id": list_id, "name": name}
    finally:
        store.close()


@app.delete("/api/lists/{list_id}")
def delete_list(request: Request, list_id: int):
    user = _require_current_user(request)
    store = get_store()
    try:
        store.delete_list(user["id"], list_id)
        return {"ok": True}
    finally:
        store.close()


@app.get("/api/lists/{list_id}/items")
def get_list_items(request: Request, list_id: int):
    user = _require_current_user(request)
    user_id = user["id"]
    store = get_store()
    try:
        search_id = _get_or_init_search_id(store, user_id)
        config = store.get_search(search_id)["config"]
        enriched = []
        for listing in store.get_list_items(user_id, list_id):
            scored = calculate_home_score(listing, config)
            purchase = calculate_purchase_estimate(listing, config)
            workflow = store.get_workflow(user_id, listing.get("source", ""), listing.get("source_listing_id", ""))
            enriched.append({**listing, "_score": scored["score"], "_components": scored["components"], "_workflow": workflow, "_purchase_estimate": purchase})
        return enriched
    finally:
        store.close()


@app.post("/api/lists/{list_id}/items")
def add_to_list(request: Request, list_id: int, body: dict):
    user = _require_current_user(request)
    source = body.get("source", "")
    sid = str(body.get("source_listing_id", ""))
    if not source or not sid:
        raise HTTPException(400, "source and source_listing_id required")
    store = get_store()
    try:
        try:
            store.add_to_list(user["id"], list_id, source, sid)
        except ValueError as exc:
            raise HTTPException(404, str(exc))
        return {"ok": True}
    finally:
        store.close()


@app.delete("/api/lists/{list_id}/items/{source}/{source_listing_id}")
def remove_from_list(request: Request, list_id: int, source: str, source_listing_id: str):
    user = _require_current_user(request)
    store = get_store()
    try:
        try:
            store.remove_from_list(user["id"], list_id, source, source_listing_id)
        except ValueError as exc:
            raise HTTPException(404, str(exc))
        return {"ok": True}
    finally:
        store.close()


@app.get("/api/notes/{source}/{source_listing_id}")
def get_note(request: Request, source: str, source_listing_id: str):
    user = _require_current_user(request)
    store = get_store()
    try:
        return {"note": store.get_note(user["id"], source, source_listing_id)}
    finally:
        store.close()


@app.put("/api/notes/{source}/{source_listing_id}")
def save_note(request: Request, source: str, source_listing_id: str, body: dict):
    user = _require_current_user(request)
    note = body.get("note", "")
    store = get_store()
    try:
        store.save_note(user["id"], source, source_listing_id, note)
        return {"ok": True}
    finally:
        store.close()


@app.post("/api/translate/selected")
def translate_selected(request: Request, body: dict):
    _require_current_user(request)
    listings = body.get("listings") or []
    if not listings:
        raise HTTPException(400, "no listings provided")

    def _worker():
        from src.buyer.collector import translate_listing
        _state["translation"] = {"total": len(listings), "done": 0, "current": ""}
        store = get_store()
        try:
            for item in listings:
                sid = str(item.get("source_listing_id", ""))
                _state["translation"]["current"] = sid
                translate_listing(store, item.get("source", ""), sid)
                _state["translation"]["done"] += 1
        finally:
            store.close()
            _state["translation"] = None

    thread = threading.Thread(target=_worker, daemon=True)
    thread.start()
    return {"ok": True, "count": len(listings)}


@app.post("/api/translate/stale")
def translate_stale(request: Request):
    _require_current_user(request)
    from src.buyer.collector import translate_listing, _needs_translation
    if _state.get("translation"):
        raise HTTPException(409, "translation already running")

    store = get_store()
    try:
        all_listings = store.latest_listings("sale") + store.latest_listings("rent")
        stale = [l for l in all_listings if _needs_translation(l)]
    finally:
        store.close()

    if not stale:
        return {"ok": True, "count": 0}

    def _worker():
        _state["translation"] = {"total": len(stale), "done": 0, "current": "", "current_address": "", "current_score": None}
        s = get_store()
        try:
            for l in stale:
                sid = str(l.get("source_listing_id", ""))
                address = l.get("address") or l.get("location") or sid
                _state["translation"]["current"] = sid
                _state["translation"]["current_address"] = address
                _state["translation"]["current_score"] = None
                translate_listing(s, l.get("source", ""), sid)
                _state["translation"]["done"] += 1
        finally:
            s.close()
            _state["translation"] = None

    thread = threading.Thread(target=_worker, daemon=True)
    thread.start()
    return {"ok": True, "count": len(stale)}


@app.get("/api/translate/stale")
def list_stale(request: Request):
    _require_current_user(request)
    from src.buyer.collector import _needs_translation
    store = get_store()
    try:
        all_listings = store.latest_listings("sale") + store.latest_listings("rent")
        stale = [
            {"source": l.get("source"), "source_listing_id": l.get("source_listing_id"), "url": l.get("url")}
            for l in all_listings if _needs_translation(l)
        ]
    finally:
        store.close()
    return {"count": len(stale), "listings": stale}


@app.get("/api/workflow/{source}/{source_listing_id}")
def get_workflow(request: Request, source: str, source_listing_id: str):
    user = _require_current_user(request)
    store = get_store()
    try:
        return store.get_workflow(user["id"], source, source_listing_id)
    finally:
        store.close()


@app.put("/api/workflow/{source}/{source_listing_id}")
def save_workflow(request: Request, source: str, source_listing_id: str, body: dict):
    user = _require_current_user(request)
    store = get_store()
    try:
        try:
            return store.save_workflow(user["id"], source, source_listing_id, body)
        except ValueError as exc:
            raise HTTPException(400, str(exc))
    finally:
        store.close()


@app.get("/api/interactions/{source}/{source_listing_id}")
def get_interactions(request: Request, source: str, source_listing_id: str):
    user = _require_current_user(request)
    store = get_store()
    try:
        return store.get_interactions(user["id"], source, source_listing_id)
    finally:
        store.close()


@app.post("/api/interactions/{source}/{source_listing_id}")
def add_interaction(request: Request, source: str, source_listing_id: str, body: dict):
    user = _require_current_user(request)
    store = get_store()
    try:
        try:
            return store.add_interaction(user["id"], source, source_listing_id, body.get("kind", ""), body.get("note", ""), body.get("occurred_at"), body.get("next_follow_up_date"))
        except ValueError as exc:
            raise HTTPException(400, str(exc))
    finally:
        store.close()


@app.delete("/api/listings/{source}/{source_listing_id}")
def delete_listing(request: Request, source: str, source_listing_id: str):
    _require_current_user(request)
    store = get_store()
    try:
        store.delete_listing(source, source_listing_id)
        return {"ok": True}
    finally:
        store.close()


@app.get("/api/changes/{source}/{source_listing_id}")
def get_listing_changes(request: Request, source: str, source_listing_id: str):
    _require_current_user(request)
    store = get_store()
    try:
        history = store.listing_history(source, source_listing_id)
        changes = []
        for index in range(1, len(history)):
            for change in diff_versions(history[index - 1]["payload"], history[index]["payload"]):
                changes.append({**change, "observed_at": history[index]["observed_at"]})
        return {"history": history, "changes": changes}
    finally:
        store.close()


@app.get("/api/duplicates")
def get_duplicates(request: Request):
    _require_current_user(request)
    store = get_store()
    try:
        return duplicate_groups(store.latest_listings("sale"))
    finally:
        store.close()


@app.post("/api/duplicates/merge")
def merge_duplicates(request: Request, body: dict):
    _require_current_user(request)
    keep = body.get("keep") or {}
    remove = body.get("remove") or []
    if not keep.get("source") or not keep.get("source_listing_id") or not isinstance(remove, list):
        raise HTTPException(400, "keep and remove are required")
    keep_key = (str(keep["source"]), str(keep["source_listing_id"]))
    remove_keys = {(str(item.get("source")), str(item.get("source_listing_id"))) for item in remove if item.get("source") and item.get("source_listing_id")}
    if keep_key in remove_keys:
        raise HTTPException(400, "the kept listing cannot also be removed")
    store = get_store()
    try:
        for source, source_listing_id in remove_keys:
            store.delete_listing(source, source_listing_id)
        return {"ok": True, "removed": len(remove_keys)}
    except Exception as exc:
        raise HTTPException(500, str(exc))
    finally:
        store.close()


@app.get("/api/alerts")
def get_alerts(request: Request):
    user = _require_current_user(request)
    user_id = user["id"]
    store = get_store()
    try:
        now = datetime.now(timezone.utc)
        for listing in store.latest_listings("sale"):
            source, sid = listing.get("source", ""), str(listing.get("source_listing_id", ""))
            history = store.listing_history(source, sid)
            if history and (now - datetime.fromisoformat(history[0]["observed_at"]).astimezone(timezone.utc)).days <= 7:
                store.add_alert(user_id, source, sid, "new", "New listing matches your search", listing.get("url"))
            for index in range(1, len(history)):
                for change in diff_versions(history[index - 1]["payload"], history[index]["payload"]):
                    if change["change_type"] in ("price_reduction", "photos_added"):
                        store.add_alert(user_id, source, sid, change["change_type"], "Price reduced" if change["change_type"] == "price_reduction" else "New photos added", listing.get("url"))
        return store.get_alerts(user_id)
    finally:
        store.close()


@app.post("/api/alerts/{alert_id}/read")
def mark_alert_read(request: Request, alert_id: int):
    user = _require_current_user(request)
    store = get_store()
    try:
        store.mark_alert_read(user["id"], alert_id)
        return {"ok": True}
    finally:
        store.close()


@app.delete("/api/alerts")
def clear_alerts(request: Request):
    user = _require_current_user(request)
    store = get_store()
    try:
        store.clear_alerts(user["id"])
        return {"ok": True}
    finally:
        store.close()


def _smart_listing_results(store, user_id, rule):
    search_id = _get_or_init_search_id(store, user_id)
    config = store.get_search(search_id)["config"]
    needs_rating = rule and rule.get("rating_min") is not None
    results = []
    for listing in store.latest_listings("sale"):
        scored = calculate_home_score(listing, config)
        purchase = calculate_purchase_estimate(listing, config)
        workflow = store.get_workflow(user_id, listing.get("source", ""), listing.get("source_listing_id", "")) if needs_rating else None
        rating = (workflow or {}).get("rating")
        if matches_rule(listing, rule, scored["score"], purchase, rating=rating):
            if not needs_rating:
                workflow = store.get_workflow(user_id, listing.get("source", ""), listing.get("source_listing_id", ""))
            results.append({**listing, "_score": scored["score"], "_components": scored["components"], "_workflow": workflow, "_purchase_estimate": purchase, "_smart_list_reason": explain_rule_match(listing, rule, scored["score"], purchase, rating=rating)})
    return results


@app.get("/api/smart-lists")
def get_smart_lists(request: Request):
    user = _require_current_user(request)
    user_id = user["id"]
    store = get_store()
    try:
        store.ensure_builtin_smart_lists(user_id, BUILTIN_SMART_LISTS)
        result = []
        for item in store.get_smart_lists(user_id):
            matches = _smart_listing_results(store, user_id, item["rule"]) if item["enabled"] else []
            result.append({**item, "item_count": len(matches)})
        return result
    finally:
        store.close()


@app.post("/api/smart-lists")
def create_smart_list(request: Request, body: dict):
    user = _require_current_user(request)
    try:
        name = (body.get("name") or "").strip()
        rule = body.get("rule") or {}
        store = get_store()
        try:
            list_id = store.create_smart_list(user["id"], name, rule)
            return {"id": list_id, "name": name, "rule": rule, "enabled": True, "is_system": False}
        finally:
            store.close()
    except sqlite3.IntegrityError:
        raise HTTPException(409, "smart list name already exists")
    except ValueError as exc:
        raise HTTPException(400, str(exc))


@app.patch("/api/smart-lists/{list_id}")
def update_smart_list(request: Request, list_id: int, body: dict):
    user = _require_current_user(request)
    user_id = user["id"]
    store = get_store()
    try:
        current = next((x for x in store.get_smart_lists(user_id) if x["id"] == list_id), None)
        if current is None:
            raise HTTPException(404, "smart list not found")
        ok = store.update_smart_list(user_id, list_id, body.get("name", current["name"]), body.get("rule", current["rule"]), body.get("enabled", current["enabled"]))
        return {"ok": ok}
    finally:
        store.close()


@app.delete("/api/smart-lists/{list_id}")
def delete_smart_list(request: Request, list_id: int):
    user = _require_current_user(request)
    store = get_store()
    try:
        try:
            if not store.delete_smart_list(user["id"], list_id):
                raise HTTPException(404, "smart list not found")
        except ValueError as exc:
            raise HTTPException(409, str(exc))
        return {"ok": True}
    finally:
        store.close()


@app.get("/api/smart-lists/{list_id}/items")
def get_smart_list_items(request: Request, list_id: int):
    user = _require_current_user(request)
    user_id = user["id"]
    store = get_store()
    try:
        current = next((x for x in store.get_smart_lists(user_id) if x["id"] == list_id), None)
        if current is None:
            raise HTTPException(404, "smart list not found")
        return _smart_listing_results(store, user_id, current["rule"]) if current["enabled"] else []
    finally:
        store.close()


EXPORT_STATUSES = {"Interested", "Contacted", "Visit planned", "Offer"}

@app.get("/api/export/interested")
def export_interested(request: Request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    user = _require_current_user(request)
    user_id = user["id"]
    store = get_store()
    try:
        search_id = _get_or_init_search_id(store, user_id)
        config = store.get_search(search_id)["config"]
        listings = store.latest_listings("sale")
        all_notes = store.get_all_notes(user_id)
        rows = []
        for listing in listings:
            src = listing.get("source", "")
            lid = str(listing.get("source_listing_id", ""))
            workflow = store.get_workflow(user_id, src, lid)
            status = (workflow or {}).get("status", "New")
            if status not in EXPORT_STATUSES:
                continue
            scored = calculate_home_score(listing, config)
            note = all_notes.get((src, lid), "")
            rows.append({
                "status": status,
                "price": listing.get("price"),
                "address": f"{listing.get('street', '')} {listing.get('house_number', '')}".strip(),
                "postcode": listing.get("postcode"),
                "city": listing.get("city"),
                "type": listing.get("property_type", ""),
                "surface_area": listing.get("surface_area"),
                "bedrooms": listing.get("bedrooms"),
                "epc": listing.get("epc_score"),
                "year": listing.get("construction_year"),
                "score": round(scored["score"]) if scored["score"] is not None else None,
                "url": listing.get("url", ""),
                "source": src,
                "agent": (workflow or {}).get("agent_name", ""),
                "agent_phone": (workflow or {}).get("agent_phone", ""),
                "agent_email": (workflow or {}).get("agent_email", ""),
                "contact_date": (workflow or {}).get("contact_date", ""),
                "follow_up": (workflow or {}).get("next_follow_up_date", ""),
                "note": note,
            })
        rows.sort(key=lambda r: (r["status"], -(r["score"] or 0)))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Interested properties"

        headers = ["Status", "Price (€)", "Address", "Postcode", "City", "Type",
                   "Surface (m²)", "Beds", "EPC", "Year", "Score", "Source",
                   "Agent", "Phone", "Email", "Contact date", "Follow-up", "Note", "URL"]
        header_fill = PatternFill("solid", fgColor="E83E8C")
        header_font = Font(bold=True, color="FFFFFF")
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        status_colors = {
            "Interested": "FEF9C3",
            "Contacted": "DBEAFE",
            "Visit planned": "D1FAE5",
            "Offer": "EDE9FE",
        }
        for row_idx, r in enumerate(rows, 2):
            values = [r["status"], r["price"], r["address"], r["postcode"], r["city"],
                      r["type"], r["surface_area"], r["bedrooms"], r["epc"], r["year"],
                      r["score"], r["source"], r["agent"], r["agent_phone"],
                      r["agent_email"], r["contact_date"], r["follow_up"], r["note"], r["url"]]
            fill_color = status_colors.get(r["status"])
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                if fill_color:
                    cell.fill = PatternFill("solid", fgColor=fill_color)

        col_widths = [14, 14, 28, 10, 16, 14, 12, 6, 6, 6, 8, 12, 18, 14, 26, 14, 12, 40, 50]
        for col, width in enumerate(col_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"interested_properties_{datetime.now().strftime('%Y%m%d')}.xlsx"
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )
    finally:
        store.close()


_dist = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "frontend", "dist"))
if os.path.isdir(_dist):
    from fastapi.staticfiles import StaticFiles
    from fastapi.responses import FileResponse

    @app.get("/{full_path:path}", include_in_schema=False)
    async def spa_fallback(full_path: str):
        file_path = os.path.join(_dist, full_path)
        if os.path.isfile(file_path):
            return FileResponse(file_path)
        return FileResponse(os.path.join(_dist, "index.html"))

    app.mount("/assets", StaticFiles(directory=os.path.join(_dist, "assets")), name="assets")
